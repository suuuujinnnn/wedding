"""Post-crawl interest analysis for the wedding community JSONL files.

The module deliberately keeps human review and external search-volume data
separate from the crawler's relevance rules.
"""

from __future__ import annotations

import csv
import hashlib
import html
import json
import math
import random
import re
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from quality_rules import analyze_post, normalize_text


ALL_SOURCE = "all"
MISSING_CLUSTER = ""
MAPPING_FIELDS = ["raw_phrase", "canonical_phrase", "category", "status", "note"]
SEARCH_MAPPING_FIELDS = [
    "canonical_topic", "community_phrase", "search_keyword", "note"
]
REVIEW_FIELDS = [
    "valid_context", "context_label", "canonical_topic", "action", "reviewer_note"
]
LABEL_FIELDS = [
    "document_id", "source", "title", "url", "predicted_relevant",
    "human_relevant", "false_positive_reason", "false_negative_reason",
    "reviewer_note",
]

# General board/UI words only. Wedding topic examples are intentionally absent.
STOPWORDS = {
    "그리고", "그러나", "그런데", "그래서", "그냥", "정말", "진짜", "너무",
    "저는", "제가", "저희", "우리", "이번", "이제", "오늘", "여러분",
    "합니다", "했습니다", "있어요", "같아요", "하는", "하고", "해서",
    "있는", "없는", "입니다", "됩니다", "되었", "이런", "그런", "어떤",
    "대한", "위한", "통해", "경우", "정도", "부분", "생각", "사람",
    "댓글", "조회", "추천검색", "비추천", "스크랩", "신고", "작성",
    "모바일에서", "앱에서", "원본", "첨부파일", "다운로드",
}
UI_PATTERNS = (
    r"\b(?:조회|추천|비추천|댓글)\s*\d+\b",
    r"\s추천검색\b.*$",
    r"원본\s+첨부파일.*$",
    r"본문\s+이미지\s+다운로드.*$",
)
TOKEN_RE = re.compile(r"[가-힣]+|[A-Za-z]+(?:[A-Za-z0-9._+-]*[A-Za-z0-9])?|\d+(?:[.,]\d+)*")
SENTENCE_RE = re.compile(r"(?<=[.!?。！？])\s+|\n+")
PARTICLES = (
    "으로부터", "에게서는", "에서는", "으로는", "이라도", "이라고", "이라는",
    "부터", "까지", "에게", "한테", "께서", "에서", "으로", "로는", "에는",
    "보다", "처럼", "만큼", "이라", "랑은", "하고", "이나", "라도",
    "은", "는", "이", "가", "을", "를", "의", "에", "와", "과", "도",
    "만", "로", "께", "랑",
)
NAVER_ALIASES = {
    "keyword": ("연관키워드", "키워드", "relKeyword", "relatedKeyword"),
    "pc_volume": ("월간PC검색수", "월간 PC 검색수", "monthlyPcQcCnt"),
    "mobile_volume": ("월간모바일검색수", "월간 모바일 검색수", "monthlyMobileQcCnt"),
    "pc_clicks": ("월평균PC클릭수", "월평균 PC 클릭수", "monthlyAvePcClkCnt"),
    "mobile_clicks": ("월평균모바일클릭수", "월평균 모바일 클릭수", "monthlyAveMobileClkCnt"),
    "competition": ("경쟁정도", "compIdx"),
    "ad_count": ("월평균노출광고수", "월평균 노출 광고 수", "plAvgDepth"),
}


@dataclass(frozen=True)
class Document:
    document_id: str
    source: str
    title: str
    body_clean: str
    body_original: str
    url: str
    created_at: str = ""
    cluster_id: str = MISSING_CLUSTER
    predicted_relevant: bool = True

    @property
    def analysis_text(self) -> str:
        return f"{self.title}. {self.body_clean}".strip()


def clean_analysis_text(text: str) -> str:
    text = html.unescape(text or "")
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"https?://\S+", " ", text)
    for pattern in UI_PATTERNS:
        text = re.sub(pattern, " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def split_sentences(text: str) -> list[str]:
    """Split without allowing an n-gram to cross a sentence boundary."""
    prepared = re.sub(r"([.!?。！？])(?=[가-힣A-Za-z0-9])", r"\1 ", text or "")
    return [part.strip() for part in SENTENCE_RE.split(prepared) if part.strip()]


def normalize_token(token: str) -> str:
    token = token.lower().strip("._+-")
    if re.fullmatch(r"[가-힣]+", token):
        for particle in PARTICLES:
            if token.endswith(particle) and len(token) - len(particle) >= 2:
                return token[:-len(particle)]
    return token


def tokenize(sentence: str) -> list[tuple[str, str]]:
    tokens = []
    for match in TOKEN_RE.finditer(sentence):
        raw = match.group(0)
        normalized = normalize_token(raw)
        if len(normalized) < 2 or normalized in STOPWORDS or normalized.isdigit():
            continue
        tokens.append((raw, normalized))
    return tokens


def sentence_ngrams(sentence: str, sizes: Iterable[int] = (2, 3)) -> list[tuple[str, str, int]]:
    tokens = tokenize(sentence)
    result = []
    for size in sizes:
        for index in range(len(tokens) - size + 1):
            window = tokens[index:index + size]
            raw = " ".join(item[0] for item in window)
            normalized = " ".join(item[1] for item in window)
            if len(set(item[1] for item in window)) == 1:
                continue
            result.append((raw, normalized, size))
    return result


def _record_signature(title: str, body: str) -> str:
    normalized = re.sub(r"[^0-9a-z가-힣]+", "", f"{title} {body}".lower())
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _similar_signature(title: str, body: str) -> str:
    normalized = re.sub(r"[^0-9a-z가-힣]+", "", f"{title} {body}".lower())
    # Reposts often only append board chrome. A long normalized prefix catches
    # those while avoiding title-only collisions.
    return normalized[:500] if len(normalized) >= 120 else normalized


def discover_input_files(input_path: Path | None, root: Path) -> list[Path]:
    if input_path:
        if input_path.is_file():
            return [input_path]
        if input_path.is_dir():
            files = sorted(input_path.glob("*.jsonl"))
            if files:
                return files
        raise FileNotFoundError(f"JSONL 입력을 찾을 수 없습니다: {input_path}")
    return [
        path for path in (root / "dc_wedding_posts.jsonl", root / "kgwed_posts.jsonl")
        if path.exists()
    ]


def load_documents(
    paths: Iterable[Path],
    source_filter: str | None = None,
    include_excluded: bool = False,
) -> tuple[list[Document], dict]:
    """Adapt the repository's real JSONL fields and remove exact/near reposts."""
    documents = []
    seen_exact: set[str] = set()
    seen_similar: set[str] = set()
    seen_titles: set[tuple[str, str]] = set()
    stats = Counter(raw_posts=0, duplicates=0, excluded=0)
    per_source = Counter()

    for path in paths:
        with path.open("r", encoding="utf-8-sig") as handle:
            for line_number, line in enumerate(handle, 1):
                if not line.strip():
                    continue
                try:
                    record = json.loads(line)
                except json.JSONDecodeError as exc:
                    warnings.warn(f"{path.name}:{line_number} JSON 오류: {exc}")
                    continue
                if record.get("record_type") == "crawl_summary":
                    continue
                stats["raw_posts"] += 1
                source = str(record.get("source") or path.stem)
                if source_filter and source != source_filter:
                    continue
                title = clean_analysis_text(str(record.get("title") or ""))
                original = str(record.get("body") or record.get("body_clean") or "")
                cleaned = clean_analysis_text(
                    str(record.get("body_clean") or record.get("clean_body") or original)
                )
                if title and cleaned.startswith(title):
                    cleaned = cleaned[len(title):].lstrip(" :-|")
                if not title and not cleaned:
                    continue
                analysis = analyze_post(title, cleaned, source)
                predicted = bool(analysis.get("keep"))
                if not predicted and not include_excluded:
                    stats["excluded"] += 1
                    continue
                exact = str(record.get("content_hash") or _record_signature(title, cleaned))
                similar = _similar_signature(title, cleaned)
                title_signature = re.sub(
                    r"[^0-9a-z가-힣]+", " ", title.lower()
                )
                title_signature = " ".join(title_signature.split())
                scoped_title = (source, title_signature)
                duplicate_title = len(title_signature) >= 12 and scoped_title in seen_titles
                if exact in seen_exact or similar in seen_similar or duplicate_title:
                    stats["duplicates"] += 1
                    continue
                seen_exact.add(exact)
                seen_similar.add(similar)
                if len(title_signature) >= 12:
                    seen_titles.add(scoped_title)
                external_id = str(
                    record.get("external_id") or record.get("post_id")
                    or record.get("id") or exact[:16]
                )
                documents.append(
                    Document(
                        document_id=external_id,
                        source=source,
                        title=title,
                        body_clean=cleaned,
                        body_original=original,
                        url=str(record.get("url") or ""),
                        created_at=str(
                            record.get("created_at") or record.get("date")
                            or record.get("written_at") or ""
                        ),
                        cluster_id=str(record.get("cluster_id") or ""),
                        predicted_relevant=predicted,
                    )
                )
                per_source[source] += 1
    stats["valid_documents"] = len(documents)
    return documents, {"counts": dict(stats), "sources": dict(per_source)}


def extract_phrases(
    documents: list[Document],
    min_document_frequency: int = 5,
    max_candidates: int = 500,
    cluster_min_documents: int = 3,
) -> list[dict]:
    scopes: list[tuple[str, list[Document]]] = [(ALL_SOURCE, documents)]
    scopes.extend(
        (source, [doc for doc in documents if doc.source == source])
        for source in sorted({doc.source for doc in documents})
    )
    rows = []
    for scope, scope_docs in scopes:
        if not scope_docs:
            continue
        occurrences = Counter()
        doc_counts = Counter()
        raw_forms: dict[str, Counter] = defaultdict(Counter)
        examples = {}
        cluster_docs: dict[tuple[str, str], set[str]] = defaultdict(set)
        for doc in scope_docs:
            found = set()
            for sentence in split_sentences(doc.analysis_text):
                for raw, normalized, size in sentence_ngrams(sentence):
                    key = (normalized, size)
                    occurrences[key] += 1
                    raw_forms[normalized][raw] += 1
                    found.add(key)
                    examples.setdefault(key, doc)
                    if doc.cluster_id:
                        cluster_docs[(normalized, doc.cluster_id)].add(doc.document_id)
            doc_counts.update(found)
        effective_min = min_document_frequency
        if len(scope_docs) < min_document_frequency:
            effective_min = max(1, min(3, len(scope_docs)))
            warnings.warn(
                f"{scope}: 유효 문서 {len(scope_docs)}건으로 최소 문서 빈도를 "
                f"{min_document_frequency}에서 {effective_min}(으)로 조정했습니다."
            )
        ranked = sorted(
            (key for key, count in doc_counts.items() if count >= effective_min),
            key=lambda key: (-doc_counts[key], -occurrences[key], key[0]),
        )[:max_candidates]
        for normalized, size in ranked:
            example = examples[(normalized, size)]
            clusters = [
                cluster for phrase, cluster in cluster_docs
                if phrase == normalized
                and len(cluster_docs[(phrase, cluster)]) >= cluster_min_documents
            ]
            rows.append(
                {
                    "phrase": raw_forms[normalized].most_common(1)[0][0],
                    "normalized_phrase": normalized,
                    "ngram_size": size,
                    "occurrence_count": occurrences[(normalized, size)],
                    "document_count": doc_counts[(normalized, size)],
                    "document_rate": round(doc_counts[(normalized, size)] / len(scope_docs), 6),
                    "source": scope,
                    "source_document_count": len(scope_docs),
                    "cluster_id": "|".join(sorted(clusters)),
                    "sample_title": example.title,
                    "sample_url": example.url,
                }
            )
    return rows


def write_csv(path: Path, rows: Iterable[dict], fieldnames: list[str] | None = None) -> None:
    rows = list(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        fieldnames = list(rows[0]) if rows else []
    if not fieldnames:
        return

    def write_to(target: Path) -> None:
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)

    try:
        write_to(path)
    except PermissionError:
        fallback = path.with_name(f"{path.stem}.updated{path.suffix}")
        warnings.warn(
            f"{path.name} 파일이 열려 있어 덮어쓰지 못했습니다. "
            f"새 결과를 {fallback.name}에 저장합니다."
        )
        write_to(fallback)


def read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def ensure_mapping(path: Path, phrase_rows: list[dict]) -> list[dict]:
    existing = read_csv(path)
    current_phrases = {
        row["normalized_phrase"]
        for row in phrase_rows if row.get("source") == ALL_SOURCE
    }
    by_raw = {}
    for row in existing:
        raw = row.get("raw_phrase", "")
        reviewed = (
            row.get("status", "").strip().lower() not in {"", "pending"}
            or bool(row.get("category", "").strip())
            or bool(row.get("note", "").strip())
            or row.get("canonical_phrase", "").strip() not in {"", raw}
        )
        if raw and (raw in current_phrases or reviewed):
            by_raw[raw] = row
    changed = not path.exists() or len(by_raw) != len(existing)
    for row in phrase_rows:
        if row.get("source") != ALL_SOURCE:
            continue
        raw = row["normalized_phrase"]
        if raw not in by_raw:
            by_raw[raw] = {
                "raw_phrase": raw,
                "canonical_phrase": raw,
                "category": "",
                "status": "pending",
                "note": "",
            }
            changed = True
    result = list(by_raw.values())
    if changed:
        write_csv(path, result, MAPPING_FIELDS)
    return result


def active_mapping(mapping_rows: list[dict]) -> dict[str, str]:
    return {
        row["raw_phrase"].strip(): (row.get("canonical_phrase") or row["raw_phrase"]).strip()
        for row in mapping_rows
        if row.get("raw_phrase", "").strip()
        and row.get("status", "").strip().lower() != "exclude"
    }


def generate_kwic(
    documents: list[Document],
    phrase_rows: list[dict],
    sample_count: int = 5,
    random_seed: int = 42,
) -> list[dict]:
    phrases = {
        row["normalized_phrase"]: row["phrase"]
        for row in phrase_rows if row.get("source") == ALL_SOURCE
    }
    rng = random.Random(random_seed)
    result = []
    for normalized, display in phrases.items():
        contexts = []
        seen = set()
        for doc in documents:
            sentences = split_sentences(doc.analysis_text)
            for index, sentence in enumerate(sentences):
                normalized_sentence = " ".join(item[1] for item in tokenize(sentence))
                if normalized not in normalized_sentence:
                    continue
                context_key = (doc.document_id, re.sub(r"\s+", " ", sentence))
                if context_key in seen:
                    continue
                seen.add(context_key)
                contexts.append(
                    {
                        "phrase": display,
                        "normalized_phrase": normalized,
                        "context": sentence,
                        "previous_sentence": sentences[index - 1] if index else "",
                        "next_sentence": sentences[index + 1] if index + 1 < len(sentences) else "",
                        "title": doc.title,
                        "source": doc.source,
                        "url": doc.url,
                        "created_at": doc.created_at,
                        "document_id": doc.document_id,
                        **{field: "" for field in REVIEW_FIELDS},
                    }
                )
        contexts.sort(key=lambda row: (row["source"], row["document_id"], row["context"]))
        if len(contexts) > sample_count:
            contexts = rng.sample(contexts, sample_count)
            contexts.sort(key=lambda row: (row["source"], row["document_id"]))
        result.extend(contexts)
    return result


def _is_trivial_relation(anchor: str, candidate: str) -> bool:
    a = anchor.replace(" ", "")
    c = candidate.replace(" ", "")
    return not c or a == c or c in a or a in c


def generate_cooccurrence(
    documents: list[Document],
    mapping_rows: list[dict],
    min_document_frequency: int = 3,
) -> list[dict]:
    mapping = active_mapping(mapping_rows)
    canonical_variants: dict[str, set[str]] = defaultdict(set)
    for raw, canonical in mapping.items():
        canonical_variants[canonical].add(raw)
    scopes = [ALL_SOURCE, *sorted({doc.source for doc in documents})]
    output = []
    for scope in scopes:
        scope_docs = documents if scope == ALL_SOURCE else [
            doc for doc in documents if doc.source == scope
        ]
        total_docs = len(scope_docs)
        anchor_docs: dict[str, set[str]] = defaultdict(set)
        candidate_docs: dict[str, set[str]] = defaultdict(set)
        pair_docs: dict[tuple[str, str], set[str]] = defaultdict(set)
        samples = {}
        for doc in scope_docs:
            doc_key = f"{doc.source}:{doc.document_id}"
            per_doc_pairs = set()
            per_doc_candidates = set()
            per_doc_anchors = set()
            for sentence in split_sentences(doc.analysis_text):
                tokens = tokenize(sentence)
                normalized_sentence = " ".join(token[1] for token in tokens)
                sentence_anchors = {
                    canonical
                    for raw, canonical in mapping.items()
                    if raw in normalized_sentence
                }
                if not sentence_anchors:
                    continue
                candidates = {token[1] for token in tokens}
                candidates.update(
                    normalized for _, normalized, _ in sentence_ngrams(sentence, (2, 3))
                )
                per_doc_candidates.update(candidates)
                per_doc_anchors.update(sentence_anchors)
                for anchor in sentence_anchors:
                    for candidate in candidates:
                        if candidate in STOPWORDS or any(
                            _is_trivial_relation(variant, candidate)
                            for variant in canonical_variants[anchor] | {anchor}
                        ):
                            continue
                        pair = (anchor, candidate)
                        per_doc_pairs.add(pair)
                        samples.setdefault(pair, (sentence, doc.url))
            for anchor in per_doc_anchors:
                anchor_docs[anchor].add(doc_key)
            for candidate in per_doc_candidates:
                candidate_docs[candidate].add(doc_key)
            for pair in per_doc_pairs:
                pair_docs[pair].add(doc_key)
        for (anchor, candidate), paired in pair_docs.items():
            count = len(paired)
            if count < min_document_frequency:
                continue
            anchor_count = len(anchor_docs[anchor])
            candidate_count = len(candidate_docs[candidate])
            if not anchor_count or not candidate_count or not total_docs:
                continue
            pxy = count / total_docs
            px = anchor_count / total_docs
            py = candidate_count / total_docs
            pmi = math.log(pxy / (px * py)) if pxy and px and py else 0.0
            npmi = pmi / -math.log(pxy) if 0 < pxy < 1 else 0.0
            context, url = samples[(anchor, candidate)]
            output.append(
                {
                    "canonical_phrase": anchor,
                    "cooccurring_phrase": candidate,
                    "co_document_count": count,
                    "anchor_document_count": anchor_count,
                    "conditional_rate": round(count / anchor_count, 6),
                    "npmi": round(npmi, 6),
                    "source": scope,
                    "sample_context": context,
                    "sample_url": url,
                }
            )
    return sorted(
        output,
        key=lambda row: (
            row["source"], row["canonical_phrase"],
            -row["co_document_count"], -row["npmi"], row["cooccurring_phrase"],
        ),
    )


def ensure_search_mapping(path: Path, mapping_rows: list[dict]) -> list[dict]:
    existing = read_csv(path)
    active_topics = set(active_mapping(mapping_rows).values())
    by_topic = {}
    for row in existing:
        topic = row.get("canonical_topic", "")
        reviewed = (
            bool(row.get("note", "").strip())
            or row.get("community_phrase", "").strip() not in {"", topic}
            or row.get("search_keyword", "").strip() not in {"", topic}
        )
        if topic and (topic in active_topics or reviewed):
            by_topic[topic] = row
    changed = not path.exists() or len(by_topic) != len(existing)
    for topic in sorted(active_topics):
        if topic not in by_topic:
            by_topic[topic] = {
                "canonical_topic": topic,
                "community_phrase": topic,
                "search_keyword": topic,
                "note": "",
            }
            changed = True
    result = list(by_topic.values())
    if changed:
        write_csv(path, result, SEARCH_MAPPING_FIELDS)
    return result


def _normalized_header(value: str) -> str:
    return re.sub(r"[\s_]+", "", str(value or "")).lower()


def _detect_column(headers: list[str], aliases: tuple[str, ...]) -> str | None:
    normalized = {_normalized_header(header): header for header in headers}
    for alias in aliases:
        if _normalized_header(alias) in normalized:
            return normalized[_normalized_header(alias)]
    return None


def _read_tabular(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        for encoding in ("utf-8-sig", "cp949"):
            try:
                with path.open("r", encoding=encoding, newline="") as handle:
                    return list(csv.DictReader(handle))
            except UnicodeDecodeError:
                continue
        raise ValueError(f"CSV 인코딩을 읽을 수 없습니다: {path}")
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            raise RuntimeError(
                "XLSX를 읽으려면 'python -m pip install openpyxl'이 필요합니다. "
                "또는 네이버 파일을 CSV로 저장해 주세요."
            ) from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value or "") for value in values[0]]
        return [dict(zip(headers, row)) for row in values[1:]]
    raise ValueError("네이버 입력은 CSV 또는 XLSX여야 합니다.")


def parse_volume(value) -> tuple[str, float | None]:
    raw = str(value if value is not None else "").strip()
    if not raw or raw == "-":
        return raw, None
    if re.fullmatch(r"<\s*10", raw):
        return raw, None
    try:
        return raw, float(raw.replace(",", ""))
    except ValueError:
        return raw, None


def import_naver(path: Path) -> list[dict]:
    raw_rows = _read_tabular(path)
    if not raw_rows:
        return []
    headers = list(raw_rows[0])
    columns = {
        name: _detect_column(headers, aliases)
        for name, aliases in NAVER_ALIASES.items()
    }
    required = [name for name in ("keyword", "pc_volume", "mobile_volume") if not columns[name]]
    if required:
        expected = ", ".join("/".join(NAVER_ALIASES[name]) for name in required)
        raise ValueError(
            f"네이버 파일의 필수 컬럼을 감지하지 못했습니다: {expected}. "
            f"실제 컬럼: {', '.join(headers)}"
        )
    result = []
    for row in raw_rows:
        keyword = str(row.get(columns["keyword"], "") or "").strip()
        if not keyword:
            continue
        pc_raw, pc_numeric = parse_volume(row.get(columns["pc_volume"]))
        mobile_raw, mobile_numeric = parse_volume(row.get(columns["mobile_volume"]))
        total = (
            pc_numeric + mobile_numeric
            if pc_numeric is not None and mobile_numeric is not None else None
        )
        result.append(
            {
                "search_keyword": keyword,
                "pc_search_volume_raw": pc_raw,
                "pc_search_volume_numeric": pc_numeric,
                "mobile_search_volume_raw": mobile_raw,
                "mobile_search_volume_numeric": mobile_numeric,
                "total_search_volume": total,
                "pc_clicks": row.get(columns["pc_clicks"], "") if columns["pc_clicks"] else "",
                "mobile_clicks": row.get(columns["mobile_clicks"], "") if columns["mobile_clicks"] else "",
                "competition": row.get(columns["competition"], "") if columns["competition"] else "",
                "average_ad_count": row.get(columns["ad_count"], "") if columns["ad_count"] else "",
            }
        )
    return result


def merge_search_volume(
    phrase_rows: list[dict],
    mapping_rows: list[dict],
    search_mapping_rows: list[dict],
    naver_rows: list[dict],
    documents: list[Document] | None = None,
) -> list[dict]:
    canonical = active_mapping(mapping_rows)
    mention_docs = Counter()
    representative = {}
    if documents is not None:
        variants = defaultdict(set)
        for raw, topic in canonical.items():
            variants[topic].add(raw)
        for doc in documents:
            text = " ".join(
                token[1]
                for sentence in split_sentences(doc.analysis_text)
                for token in tokenize(sentence)
            )
            for topic, raw_phrases in variants.items():
                if any(raw in text for raw in raw_phrases):
                    mention_docs[topic] += 1
    for row in phrase_rows:
        if row.get("source") == ALL_SOURCE:
            topic = canonical.get(row["normalized_phrase"])
            if topic:
                representative.setdefault(topic, row["phrase"])
    naver_by_keyword = {
        _normalized_header(row["search_keyword"]): row for row in naver_rows
    }
    output = []
    for item in search_mapping_rows:
        topic = item.get("canonical_topic", "").strip()
        keyword = item.get("search_keyword", "").strip()
        naver = naver_by_keyword.get(_normalized_header(keyword), {})
        output.append(
            {
                "canonical_topic": topic,
                "community_phrase": item.get("community_phrase") or representative.get(topic, topic),
                "community_document_count": mention_docs.get(topic, 0),
                "search_keyword": keyword,
                **naver,
            }
        )
    return output


def _median(values: list[float]) -> float | None:
    values = sorted(values)
    if not values:
        return None
    middle = len(values) // 2
    return (
        values[middle] if len(values) % 2
        else (values[middle - 1] + values[middle]) / 2
    )


def build_final_topics(
    documents: list[Document],
    phrase_rows: list[dict],
    mapping_rows: list[dict],
    kwic_rows: list[dict],
    cooccurrence_rows: list[dict],
    merged_search_rows: list[dict],
    community_high_threshold: float | None = None,
    search_high_threshold: float | None = None,
) -> list[dict]:
    canonical = active_mapping(mapping_rows)
    topic_docs: dict[str, set[str]] = defaultdict(set)
    source_docs: dict[str, Counter] = defaultdict(Counter)
    sample_url = {}
    representative = {}
    normalized_topics = defaultdict(set)
    for raw, topic in canonical.items():
        normalized_topics[topic].add(raw)
    for doc in documents:
        doc_key = f"{doc.source}:{doc.document_id}"
        text = " ".join(token[1] for sentence in split_sentences(doc.analysis_text) for token in tokenize(sentence))
        for topic, variants in normalized_topics.items():
            if any(variant in text for variant in variants):
                topic_docs[topic].add(doc_key)
                source_docs[topic][doc.source] += 1
                sample_url.setdefault(topic, doc.url)
    for row in phrase_rows:
        if row.get("source") == ALL_SOURCE:
            topic = canonical.get(row["normalized_phrase"])
            if topic:
                representative.setdefault(topic, row["phrase"])
    reviewed = defaultdict(lambda: [0, 0])
    for row in kwic_rows:
        topic = row.get("canonical_topic") or canonical.get(row.get("normalized_phrase", ""), "")
        value = str(row.get("valid_context", "")).strip().lower()
        if topic and value:
            reviewed[topic][1] += 1
            if value in {"1", "true", "y", "yes", "예", "유효"}:
                reviewed[topic][0] += 1
    co_top = defaultdict(list)
    for row in cooccurrence_rows:
        if row.get("source") == ALL_SOURCE and len(co_top[row["canonical_phrase"]]) < 5:
            co_top[row["canonical_phrase"]].append(row["cooccurring_phrase"])
    search_by_topic = {row["canonical_topic"]: row for row in merged_search_rows}
    topics = sorted(topic_docs)
    community_values = [float(len(topic_docs[topic])) for topic in topics]
    search_values = [
        float(search_by_topic[topic]["total_search_volume"])
        for topic in topics
        if search_by_topic.get(topic, {}).get("total_search_volume") not in (None, "")
    ]
    community_threshold = (
        community_high_threshold
        if community_high_threshold is not None else _median(community_values)
    )
    search_threshold = (
        search_high_threshold
        if search_high_threshold is not None else _median(search_values)
    )
    enough_quadrants = len(topics) >= 4 and len(search_values) >= 4
    rows = []
    for topic in topics:
        search = search_by_topic.get(topic, {})
        community_count = len(topic_docs[topic])
        total_search = search.get("total_search_volume")
        if not enough_quadrants or total_search in (None, ""):
            quadrant = "insufficient_data"
            directions = ("사람 검수 후 정보성 소재 후보", "KWIC로 문제 원인 확인", "이해관계자 연결 검토")
        else:
            community_high = community_count >= community_threshold
            search_high = float(total_search) >= search_threshold
            if community_high and search_high:
                quadrant = "community_high_search_high"
                directions = ("우선 제작할 정보성 콘텐츠 후보", "반복 질문 구체화", "주요 이해관계자 우선 조사")
            elif community_high:
                quadrant = "community_high_search_low"
                directions = ("커뮤니티 잠재 문제 설명", "원인 확인 인터뷰 우선", "마찰 지점 조사")
            elif search_high:
                quadrant = "community_low_search_high"
                directions = ("검색 유입형 콘텐츠 후보", "누락 출처 여부 질문", "수집 범위 누락 점검")
            else:
                quadrant = "community_low_search_low"
                directions = ("후순위·추가 검증", "필요 시 탐색 질문", "추가 표본 후 판단")
        valid, reviewed_count = reviewed[topic]
        rows.append(
            {
                "canonical_topic": topic,
                "community_phrase": representative.get(topic, topic),
                "community_document_count": community_count,
                "mentions_per_1000_valid_documents": round(community_count / len(documents) * 1000, 3) if documents else "",
                "source_document_counts": json.dumps(source_docs[topic], ensure_ascii=False),
                "source_diversity": len(source_docs[topic]),
                "kwic_review_pass_rate": round(valid / reviewed_count, 6) if reviewed_count else "",
                "kwic_reviewed_count": reviewed_count,
                "top_cooccurring_phrases": " | ".join(co_top[topic]),
                "search_keyword": search.get("search_keyword", ""),
                "mobile_search_volume": search.get("mobile_search_volume_raw", ""),
                "pc_search_volume": search.get("pc_search_volume_raw", ""),
                "total_search_volume": total_search if total_search is not None else "",
                "classification": quadrant,
                "content_direction": directions[0],
                "interview_direction": directions[1],
                "ecosystem_direction": directions[2],
                "sample_url": sample_url.get(topic, ""),
            }
        )
    return rows


def build_label_template(documents: list[Document]) -> list[dict]:
    return [
        {
            "document_id": doc.document_id,
            "source": doc.source,
            "title": doc.title,
            "url": doc.url,
            "predicted_relevant": str(doc.predicted_relevant).lower(),
            "human_relevant": "",
            "false_positive_reason": "",
            "false_negative_reason": "",
            "reviewer_note": "",
        }
        for doc in documents
    ]


def detect_source_templates(documents: list[Document], min_documents: int = 3) -> list[dict]:
    """Surface repeated board/UI/boilerplate text for review, not auto-filtering."""
    counts: dict[tuple[str, str], set[str]] = defaultdict(set)
    display = {}
    template_type = {}
    for doc in documents:
        doc_key = f"{doc.source}:{doc.document_id}"
        original = normalize_text(doc.body_original)
        ui_match = re.search(r"추천검색\b.*?(?:스크랩\s+신고|신고)", original)
        if ui_match:
            normalized = re.sub(r"\d+", "#", ui_match.group(0))
            normalized = re.sub(r"\s+", " ", normalized).strip()
            key = (doc.source, normalized)
            counts[key].add(doc_key)
            display[key] = ui_match.group(0)[:300]
            template_type[key] = "board_ui"
        seen_sentences = set()
        for sentence in split_sentences(original):
            normalized = re.sub(r"\d+", "#", sentence)
            normalized = re.sub(r"\s+", " ", normalized).strip()
            if not 15 <= len(normalized) <= 220 or normalized in seen_sentences:
                continue
            seen_sentences.add(normalized)
            key = (doc.source, normalized)
            counts[key].add(doc_key)
            display.setdefault(key, sentence[:300])
            template_type.setdefault(key, "repeated_boilerplate")
    rows = [
        {
            "source": source,
            "template_text": display[(source, normalized)],
            "template_type": template_type[(source, normalized)],
            "document_count": len(documents_found),
            "reviewer_note": "",
        }
        for (source, normalized), documents_found in counts.items()
        if len(documents_found) >= min_documents
    ]
    return sorted(rows, key=lambda row: (row["source"], -row["document_count"], row["template_text"]))


def write_run_metadata(path: Path, stats: dict, warnings_list: list[str]) -> None:
    payload = {
        "input_fields": {
            "title": "title",
            "clean_body": "body_clean (fallback: clean_body/body)",
            "source": "source",
            "url": "url",
            "document_id": "external_id (fallback: post_id/id/content hash)",
            "created_at": "created_at/date/written_at; current files do not provide it",
            "cluster_id": "cluster_id; current files do not provide it",
        },
        "stats": stats,
        "warnings": warnings_list,
        "interpretation_limits": [
            "커뮤니티 언급률은 전체 예비부부의 비율이 아닙니다.",
            "업체 운영 후기 게시판과 일반 커뮤니티의 성격이 다릅니다.",
            "사람이 KWIC를 검수하기 전에는 후보를 확정 관심사로 해석하지 않습니다.",
            "네이버 검색량은 인스타그램 내부 검색량이 아닙니다.",
        ],
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
